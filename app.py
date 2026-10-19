import io
import copy
import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import column_index_from_string
import pandas as pd
import streamlit as st

st.set_page_config(
    page_title="Grelha de Avaliação por Formando - ENB",
    page_icon="📋",
    layout="centered"
)

st.title("📋 Gerador de Grelha de Avaliação (1 Tab por Formando)")
st.write("Carregue a grelha modelo e a lista de formandos para gerar um livro Excel com um separador individualizado para cada formando e a pauta geral preenchida.")

def definir_valor_celula_seguro(sheet, row, col, value):
    """Escreve o valor na célula tratando com segurança células mescladas (MergedCells)."""
    cell = sheet.cell(row=row, column=col)
    if isinstance(cell, MergedCell):
        for rng in sheet.merged_cells.ranges:
            if cell.coordinate in rng:
                top_left_cell = sheet.cell(row=rng.min_row, column=rng.min_col)
                top_left_cell.value = value
                return
    else:
        cell.value = value

def extrair_nomes_formandos(file_formandos, coluna_letra="K"):
    """Extrai os nomes dos formandos a partir da coluna especificada."""
    col_idx = column_index_from_string(coluna_letra.upper()) - 1

    if file_formandos.name.endswith(".csv"):
        df = pd.read_csv(file_formandos, header=None)
    else:
        df = pd.read_excel(file_formandos, header=None)

    if df.shape[1] <= col_idx:
        raise ValueError(f"O ficheiro possui apenas {df.shape[1]} colunas, mas a Coluna {coluna_letra.upper()} foi selecionada.")

    coluna_dados = df.iloc[:, col_idx]
    palavras_ignorar = ["nome", "nome do formando", "formando", "nome completo", "formandos", "nome:", "cform", "tipo", "data"]

    nomes = []
    for val in coluna_dados:
        if pd.isna(val) or val is None:
            continue
        val_str = str(val).strip()
        if not val_str or val_str.lower() in palavras_ignorar or any(val_str.lower().startswith(p) for p in palavras_ignorar):
            continue
        nomes.append(val_str)

    return nomes

def gerar_grelha_com_tabs_individuais(file_avaliacao, file_formandos, coluna_letra="K"):
    """Gera um workbook Excel com a pauta geral e um separador (tab) por formando."""
    nomes_formandos = extrair_nomes_formandos(file_formandos, coluna_letra)

    if not nomes_formandos:
        raise ValueError(f"Nenhum nome válido encontrado na Coluna {coluna_letra.upper()}.")

    wb = openpyxl.load_workbook(file_avaliacao)
    aba_modelo = wb.active

    # 1. Preencher a Pauta Geral (Aba principal)
    linha_inicial = 14
    for idx, nome in enumerate(nomes_formandos):
        linha_atual = linha_inicial + idx
        definir_valor_celula_seguro(aba_modelo, row=linha_atual, col=2, value=idx + 1)
        definir_valor_celula_seguro(aba_modelo, row=linha_atual, col=11, value=nome)

    # 2. Criar um separador (tab) individual por formando
    for idx, nome in enumerate(nomes_formandos):
        # Limpar o nome da aba para caracteres válidos no Excel (máx 31 caracteres)
        nome_aba = f"{idx+1}_{nome[:25]}".translate(str.maketrans("", "", r"/*?:[]\\"))
        
        # Duplicar a folha modelo
        aba_formando = wb.copy_worksheet(aba_modelo)
        aba_formando.title = nome_aba

        # Preencher o cabeçalho individual do formando na respetiva aba
        definir_valor_celula_seguro(aba_formando, row=14, col=2, value=idx + 1)
        definir_valor_celula_seguro(aba_formando, row=14, col=11, value=nome)

    output_buffer = io.BytesIO()
    wb.save(output_buffer)
    output_buffer.seek(0)

    return output_buffer, len(nomes_formandos)

# --- INTERFACE STREAMLIT ---

col1, col2 = st.columns(2)

with col1:
    st.subheader("1. Grelha Modelo")
    file_eval = st.file_uploader("Carregue a Grelha (.xlsx)", type=["xlsx"], key="eval")

with col2:
    st.subheader("2. Ficheiro Formandos")
    file_list = st.file_uploader("Carregue os Formandos (.xlsx, .csv)", type=["xlsx", "xls", "csv"], key="list")

col_opcao = st.text_input("Letra da Coluna com os Nomes:", value="K", max_chars=2).upper()

if file_eval and file_list:
    if st.button("🚀 Gerar Grelha com Separadores Individuais", type="primary"):
        try:
            with st.spinner("A duplicar separadores e a preencher os campos..."):
                excel_gerado, total = gerar_grelha_com_tabs_individuais(file_eval, file_list, coluna_letra=col_opcao)

            st.success(f"✅ Sucesso! Foram criados {total} separadores individuais e preenchida a pauta geral.")
            
            st.download_button(
                label="📥 Descarregar Ficheiro de Avaliação Completo",
                data=excel_gerado,
                file_name="Grelha_Avaliacao_Formandos_Tabs.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
        except Exception as e:
            st.error(f"Erro ao processar: {e}")
